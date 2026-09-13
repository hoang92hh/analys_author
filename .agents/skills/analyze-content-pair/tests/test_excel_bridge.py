from __future__ import annotations

import copy
import importlib.util
import json
import sys
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from openpyxl import Workbook, load_workbook

SCRIPT = Path(__file__).parents[1] / "scripts" / "excel_bridge.py"
SPEC = importlib.util.spec_from_file_location("excel_bridge", SCRIPT)
bridge = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
sys.modules[SPEC.name] = bridge
SPEC.loader.exec_module(bridge)


class ExcelBridgeTests(unittest.TestCase):
    def setUp(self):
        # Fixtures and generated workbooks stay inside the project workspace.
        self.temp = tempfile.TemporaryDirectory(dir=Path.cwd())
        self.root = Path(self.temp.name)
        (self.root / "reference").mkdir()
        (self.root / "criteria").mkdir()
        (self.root / "reference/a.txt").write_text(
            "- Part1--\nNguồn A một\n\n--Part2--\nNguồn A hai\n", encoding="utf-8"
        )
        # Reversed source order ensures pairing is based on ID rather than position.
        (self.root / "reference/b.txt").write_text(
            "-- Part2--\nNội dung B hai\n\n--Part1--\nNội dung B một\n", encoding="utf-8"
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Criteria"
        sheet.append(["ID", "Nhóm", "Tiêu chí", "Requirement", "Analysis Instruction", "Result"])
        sheet.append(["C002", "G2", "Second", "Analyze second", "Compare selection", None])
        sheet.append(["C001", "G1", "First", "Analyze first", "Inspect sentences", None])
        notes = workbook.create_sheet("Notes")
        notes["A1"], notes["B1"] = "Keep me", "=1+1"
        workbook.save(self.root / "criteria/source.xlsx")
        workbook.close()
        config = self.root / "config.yaml"
        config.write_text("\n".join([
            "content_a: reference/a.txt", "content_b: reference/b.txt",
            "criteria: criteria/source.xlsx", "criteria_payload: output/payload.json",
            "analysis_results: output/results.json", "output: output/result_analysis.xlsx",
            "result_column: Result", "expected_criteria_count: 2",
        ]), encoding="utf-8")
        self.paths = bridge.read_config(config, self.root)

    def tearDown(self):
        self.temp.cleanup()

    def _document(self, payload):
        return {
            "schema_version": bridge.SCHEMA_VERSION,
            "payload_digest": payload["payload_digest"],
            "parts": [
                {"part_id": part["part_id"], "results": [
                    {"criterion_key": criterion["criterion_key"],
                     "criterion_id": criterion["criterion_id"],
                     "result": f"Only {part['part_id']} / {criterion['criterion_id']}"}
                    for criterion in payload["criteria"]
                ]}
                for part in payload["parts"]
            ],
        }

    def _write(self, document):
        self.paths.analysis_results.parent.mkdir(parents=True, exist_ok=True)
        self.paths.analysis_results.write_text(json.dumps(document, ensure_ascii=False), encoding="utf-8")

    def test_prepare_pairs_by_id_and_preserves_criteria_order_and_instructions(self):
        payload = bridge.prepare(self.paths)
        self.assertEqual([p["part_id"] for p in payload["parts"]], ["Part1", "Part2"])
        self.assertEqual(payload["parts"][0]["content_b"], "Nội dung B một")
        self.assertEqual(payload["parts"][1]["content_a"], "Nguồn A hai")
        self.assertEqual([c["criterion_key"] for c in payload["criteria"]],
                         ["Criteria!2:C002", "Criteria!3:C001"])
        self.assertEqual(payload["criteria"][0]["requirement"], "Analyze second")
        self.assertEqual(payload["criteria"][0]["analysis_instruction"], "Compare selection")
        self.assertEqual(payload["payload_digest"], bridge.payload_digest(payload))
        self.assertFalse(self.paths.output.exists())

    def test_finalize_keeps_independent_results_in_part_sheets_and_preserves_sources(self):
        sources = [self.paths.content_a, self.paths.content_b, self.paths.criteria]
        before = {p: p.read_bytes() for p in sources}
        payload = bridge.prepare(self.paths)
        self._write(self._document(payload))
        self.assertEqual(bridge.finalize(self.paths), 4)
        output = load_workbook(self.paths.output, data_only=False)
        try:
            self.assertEqual(output.sheetnames, ["Part1", "Part2"])
            for name in output.sheetnames:
                rows = list(output[name].iter_rows(values_only=True))
                self.assertEqual(rows, [
                    ("ID", "Nhóm", "Tiêu chí", "Requirement", "Result"),
                    ("C002", "G2", "Second", "Analyze second", f"Only {name} / C002"),
                    ("C001", "G1", "First", "Analyze first", f"Only {name} / C001"),
                ])
        finally:
            output.close()
        self.assertEqual(before, {p: p.read_bytes() for p in sources})

    def test_empty_results_remain_blank_without_dropping_any_criterion(self):
        payload = bridge.prepare(self.paths)
        document = self._document(payload)
        document["parts"][0]["results"][0]["result"] = ""
        document["parts"][0]["results"][1]["result"] = "   "
        self._write(document)
        self.assertEqual(bridge.finalize(self.paths), 4)
        output = load_workbook(self.paths.output)
        try:
            self.assertEqual(output["Part1"].max_row, 3)
            self.assertEqual([output["Part1"].cell(r, 1).value for r in (2, 3)], ["C002", "C001"])
            self.assertIsNone(output["Part1"]["E2"].value)
            self.assertIsNone(output["Part1"]["E3"].value)
            self.assertEqual(output["Part2"]["E2"].value, "Only Part2 / C002")
        finally:
            output.close()

    def test_missing_duplicate_unknown_and_out_of_order_parts_are_rejected(self):
        payload = bridge.prepare(self.paths)
        original = self._document(payload)
        variants = [
            original["parts"][:1],
            [original["parts"][0], original["parts"][0]],
            list(reversed(original["parts"])),
            [original["parts"][0], {"part_id": "Part9", "results": []}],
        ]
        for parts in variants:
            with self.subTest(parts=parts):
                self._write({**original, "parts": parts})
                with self.assertRaises(bridge.BridgeError):
                    bridge.finalize(self.paths)
                self.assertFalse(self.paths.output.exists())

    def test_invalid_criterion_documents_are_rejected(self):
        payload = bridge.prepare(self.paths)
        original = self._document(payload)
        variants = []
        for change in ("missing", "duplicate", "unknown", "order", "id", "type", "field", "length"):
            document = copy.deepcopy(original)
            results = document["parts"][0]["results"]
            if change == "missing":
                results.pop()
            elif change == "duplicate":
                results[1] = copy.deepcopy(results[0])
            elif change == "unknown":
                results[0]["criterion_key"] = "Criteria!2:UNKNOWN"
            elif change == "order":
                results.reverse()
            elif change == "id":
                results[0]["criterion_id"] = "C999"
            elif change == "type":
                results[0]["result"] = None
            elif change == "field":
                results[0]["frequency"] = 1
            else:
                results[0]["result"] = "x" * 32768
            variants.append((change, document))
        for change, document in variants:
            with self.subTest(change=change):
                self._write(document)
                with self.assertRaises(bridge.BridgeError):
                    bridge.finalize(self.paths)
                self.assertFalse(self.paths.output.exists())

    def test_changed_sources_and_edited_payload_are_rejected(self):
        for target in ("content_a", "content_b", "criteria", "criteria_payload"):
            with self.subTest(target=target):
                payload = bridge.prepare(self.paths)
                self._write(self._document(payload))
                path = getattr(self.paths, target)
                before = path.read_bytes()
                try:
                    if target == "criteria":
                        workbook = load_workbook(path)
                        workbook["Criteria"]["E2"] = "Changed instruction"
                        workbook.save(path)
                        workbook.close()
                    elif target == "criteria_payload":
                        payload["criteria"][0]["row"] = 99
                        path.write_text(json.dumps(payload), encoding="utf-8")
                    else:
                        path.write_bytes(before + b"\nChanged source")
                    with self.assertRaisesRegex(bridge.BridgeError, "stale"):
                        bridge.finalize(self.paths)
                    self.assertFalse(self.paths.output.exists())
                finally:
                    path.write_bytes(before)

    def test_old_results_cannot_be_reused_after_prepare_runs_again(self):
        payload = bridge.prepare(self.paths)
        self._write(self._document(payload))
        self.paths.content_a.write_bytes(self.paths.content_a.read_bytes() + b"\nChanged source")
        bridge.prepare(self.paths)
        with self.assertRaisesRegex(bridge.BridgeError, "different payload"):
            bridge.finalize(self.paths)

    def test_invalid_source_parts_fail_before_payload_is_written(self):
        invalid = [
            "No markers", "--Part1--\nA\n-- Part1--\nDuplicate",
            "--Part1--\nA\n--Part2--\n", "Preamble\n--Part1--\nA",
            "--Part1--\nA", "--Part0--\nA",
        ]
        for text in invalid:
            with self.subTest(text=text):
                self.paths.content_a.write_text(text, encoding="utf-8")
                with self.assertRaises(bridge.BridgeError):
                    bridge.prepare(self.paths)
                self.assertFalse(self.paths.criteria_payload.exists())

    def test_missing_analysis_instruction_column_and_duplicate_ids_are_rejected(self):
        before = self.paths.criteria.read_bytes()
        for change in ("header", "duplicate", "no_id"):
            with self.subTest(change=change):
                workbook = load_workbook(self.paths.criteria)
                sheet = workbook["Criteria"]
                if change == "header":
                    sheet["E1"] = "Wrong header"
                elif change == "duplicate":
                    sheet["A3"] = "C002"
                else:
                    sheet["A3"] = None
                workbook.save(self.paths.criteria)
                workbook.close()
                try:
                    with self.assertRaises(bridge.BridgeError):
                        bridge.prepare(self.paths)
                finally:
                    self.paths.criteria.write_bytes(before)

    def test_default_count_is_22_and_mismatching_master_is_rejected(self):
        config = self.root / "config.yaml"
        config.write_text(config.read_text(encoding="utf-8").replace("expected_criteria_count: 2", ""), encoding="utf-8")
        paths = bridge.read_config(config, self.root)
        self.assertEqual(paths.expected_criteria_count, 22)
        with self.assertRaisesRegex(bridge.BridgeError, "Expected 22 criteria, found 2"):
            bridge.prepare(paths)

    def test_output_replacement_requires_intent_and_failed_validation_preserves_output(self):
        payload = bridge.prepare(self.paths)
        document = self._document(payload)
        self._write(document)
        bridge.finalize(self.paths)
        before = self.paths.output.read_bytes()
        with self.assertRaisesRegex(bridge.BridgeError, "already exists"):
            bridge.finalize(self.paths)
        document["parts"].pop()
        self._write(document)
        with self.assertRaises(bridge.BridgeError):
            bridge.finalize(self.paths, overwrite=True)
        self.assertEqual(self.paths.output.read_bytes(), before)
        document = self._document(payload)
        document["parts"][0]["results"][0]["result"] = "Changed Result"
        self._write(document)
        bridge.finalize(self.paths, overwrite=True)
        output = load_workbook(self.paths.output)
        try:
            self.assertEqual(output["Part1"]["E2"].value, "Changed Result")
        finally:
            output.close()

    def test_formula_like_result_is_exported_as_text(self):
        payload = bridge.prepare(self.paths)
        document = self._document(payload)
        document["parts"][0]["results"][0]["result"] = "=Observed wording"
        self._write(document)
        bridge.finalize(self.paths)
        output = load_workbook(self.paths.output, data_only=False)
        try:
            self.assertEqual(output["Part1"]["E2"].value, "=Observed wording")
            self.assertEqual(output["Part1"]["E2"].data_type, "s")
        finally:
            output.close()

    def test_overwriting_sources_and_path_collisions_are_rejected(self):
        payload = bridge.prepare(self.paths)
        self._write(self._document(payload))
        for source in (self.paths.content_a, self.paths.content_b, self.paths.criteria):
            with self.subTest(source=source):
                with self.assertRaisesRegex(bridge.BridgeError, "input file"):
                    bridge.finalize(replace(self.paths, output=source), overwrite=True)
        config = self.root / "config.yaml"
        config.write_text(config.read_text(encoding="utf-8").replace(
            "output: output/result_analysis.xlsx", "output: reference/a.txt"), encoding="utf-8")
        with self.assertRaisesRegex(bridge.BridgeError, "distinct"):
            bridge.read_config(config, self.root)

    def test_duplicate_json_fields_are_rejected(self):
        self.paths.analysis_results.parent.mkdir(parents=True, exist_ok=True)
        self.paths.analysis_results.write_text('{"parts": [], "parts": []}', encoding="utf-8")
        with self.assertRaisesRegex(bridge.BridgeError, "Duplicate JSON field"):
            bridge._load_json(self.paths.analysis_results)

    def test_numeric_part_order_and_bom_are_supported(self):
        for path in (self.paths.content_a, self.paths.content_b):
            path.write_text("--Part10--\nTen\n-- Part2--\nTwo\n--part1--\nOne", encoding="utf-8-sig")
        payload = bridge.prepare(self.paths)
        self.assertEqual([p["part_id"] for p in payload["parts"]], ["Part1", "Part2", "Part10"])

    def test_real_project_inputs_export_five_sheets_with_22_criteria_and_blank_cells(self):
        # Transport smoke test only: synthetic Result text is not content analysis.
        root = Path.cwd()
        paths = bridge.read_config(root / ".agents/skills/analyze-content-pair/config.yaml", root)
        paths = replace(paths, criteria_payload=self.root / "real/payload.json",
                        analysis_results=self.root / "real/results.json", output=self.root / "real/result_analysis.xlsx")
        before = {p: p.read_bytes() for p in (paths.content_a, paths.content_b, paths.criteria)}
        payload = bridge.prepare(paths)
        self.assertEqual(len(payload["criteria"]), 22)
        self.assertEqual([p["part_id"] for p in payload["parts"]], [f"Part{i}" for i in range(1, 6)])
        document = self._document(payload)
        for part in document["parts"]:
            part["results"][0]["result"] = ""
        paths.analysis_results.write_text(json.dumps(document, ensure_ascii=False), encoding="utf-8")
        self.assertEqual(bridge.finalize(paths), 110)
        output = load_workbook(paths.output)
        try:
            self.assertEqual(output.sheetnames, [f"Part{i}" for i in range(1, 6)])
            for sheet in output.worksheets:
                self.assertEqual(sheet.max_row, 23)
                self.assertEqual(sheet.max_column, 5)
                self.assertEqual([c.value for c in sheet[1]], bridge.OUTPUT_HEADERS)
                self.assertIsNone(sheet["E2"].value)
                self.assertEqual([sheet.cell(r, 1).value for r in range(2, 24)],
                                 [c["criterion_id"] for c in payload["criteria"]])
                self.assertEqual(sheet["E3"].value, f"Only {sheet.title} / {payload['criteria'][1]['criterion_id']}")
        finally:
            output.close()
        self.assertEqual(before, {p: p.read_bytes() for p in before})


if __name__ == "__main__":
    unittest.main()
