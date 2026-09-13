#!/usr/bin/env python3
"""Excel I/O for independent A.Part_i -> B.Part_i analyses."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Sequence, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SCHEMA_VERSION = "3.0"
OUTPUT_HEADERS = ["ID", "Nhóm", "Tiêu chí", "Requirement", "Result"]
REQUIRED_CONFIG = (
    "content_a", "content_b", "criteria", "criteria_payload",
    "analysis_results", "output", "result_column",
)
PART_MARKER = re.compile(r"^[ \t]*-+[ \t]*Part[ \t]*([0-9]+)[ \t]*-+[ \t]*\r?$", re.MULTILINE | re.IGNORECASE)


class BridgeError(RuntimeError):
    pass


@dataclass(frozen=True)
class Paths:
    project_root: Path
    content_a: Path
    content_b: Path
    criteria: Path
    criteria_payload: Path
    analysis_results: Path
    output: Path
    result_column: str
    expected_criteria_count: int = 22


def read_config(config_path: Path, project_root: Path | None = None) -> Paths:
    values: Dict[str, str] = {}
    for number, raw in enumerate(config_path.resolve().read_text(encoding="utf-8-sig").splitlines(), 1):
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        if ":" not in raw:
            raise BridgeError(f"Config line {number} is not key: value")
        key, value = raw.split(":", 1)
        key, value = key.strip(), value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in {'"', "'"}:
            value = value[1:-1]
        if not key or key in values:
            raise BridgeError(f"Invalid or duplicate config key at line {number}")
        values[key] = value
    missing = [key for key in REQUIRED_CONFIG if not values.get(key)]
    if missing:
        raise BridgeError("Missing config values: " + ", ".join(missing))
    if values["result_column"] != "Result":
        raise BridgeError("result_column must be Result for the five-column output")
    try:
        count = int(values.get("expected_criteria_count", "22"))
    except ValueError as error:
        raise BridgeError("expected_criteria_count must be a positive integer") from error
    if count <= 0:
        raise BridgeError("expected_criteria_count must be a positive integer")
    root = (project_root or Path.cwd()).resolve()

    def resolve(value: str) -> Path:
        path = Path(value)
        return (path if path.is_absolute() else root / path).resolve()

    resolved = [resolve(values[key]) for key in REQUIRED_CONFIG[:-1]]
    if len(set(resolved)) != len(resolved):
        raise BridgeError("Input and output paths must all be distinct")
    return Paths(root, *resolved, values["result_column"], count)


def _require_files(paths: Iterable[Path]) -> None:
    missing = [str(path) for path in paths if not path.is_file()]
    if missing:
        raise BridgeError("Missing input file(s): " + ", ".join(missing))


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def payload_digest(payload: Mapping[str, Any]) -> str:
    body = {key: value for key, value in payload.items() if key != "payload_digest"}
    encoded = json.dumps(body, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _headers(sheet: Any) -> Dict[str, int]:
    headers: Dict[str, int] = {}
    for cell in sheet[1]:
        if cell.value is None:
            continue
        name = str(cell.value).strip()
        if name in headers:
            raise BridgeError(f"Duplicate header {name!r} in sheet {sheet.title!r}")
        headers[name] = cell.column
    return headers


def _relative(path: Path, root: Path) -> str:
    try:
        return path.relative_to(root).as_posix()
    except ValueError:
        return str(path)


def split_parts(text: str, label: str) -> Dict[str, str]:
    markers = list(PART_MARKER.finditer(text))
    if not markers:
        raise BridgeError(f"{label}: no Part markers found")
    if text[:markers[0].start()].strip():
        raise BridgeError(f"{label}: content before the first Part marker")
    parts: Dict[str, str] = {}
    for index, marker in enumerate(markers):
        number = int(marker.group(1))
        part_id = f"Part{number}"
        if number <= 0 or len(part_id) > 31:
            raise BridgeError(f"{label}: invalid Part ID {marker.group(1)!r}")
        if part_id in parts:
            raise BridgeError(f"{label}: duplicate Part ID {part_id}")
        end = markers[index + 1].start() if index + 1 < len(markers) else len(text)
        content = text[marker.end():end].strip()
        if not content:
            raise BridgeError(f"{label}: empty {part_id}")
        parts[part_id] = content
    return parts


def extract_criteria(paths: Paths) -> Dict[str, Any]:
    """Read criteria and pair sources without writing to any source file."""
    _require_files((paths.content_a, paths.content_b, paths.criteria))
    a_parts = split_parts(paths.content_a.read_text(encoding="utf-8-sig"), "Content A")
    b_parts = split_parts(paths.content_b.read_text(encoding="utf-8-sig"), "Content B")
    if set(a_parts) != set(b_parts):
        missing_b = sorted(set(a_parts) - set(b_parts))
        missing_a = sorted(set(b_parts) - set(a_parts))
        raise BridgeError(f"Unpaired Parts; missing in B: {missing_b}; missing in A: {missing_a}")
    parts = [
        {"part_id": part_id, "content_a": a_parts[part_id], "content_b": b_parts[part_id]}
        for part_id in sorted(a_parts, key=lambda name: int(name[4:]))
    ]
    workbook = load_workbook(paths.criteria, read_only=False, data_only=False)
    criteria: List[Dict[str, Any]] = []
    seen_ids = set()
    criteria_sheets = 0
    required = ("ID", "Nhóm", "Tiêu chí", "Requirement", "Analysis Instruction")
    try:
        for sheet in workbook.worksheets:
            headers = _headers(sheet)
            present = [name in headers for name in required]
            if any(present) and not all(present):
                absent = [name for name in required if name not in headers]
                raise BridgeError(f"Sheet {sheet.title!r} has partial criteria schema; missing: {', '.join(absent)}")
            if not all(present):
                continue
            criteria_sheets += 1
            if paths.result_column not in headers:
                raise BridgeError(f"Sheet {sheet.title!r} is missing result column {paths.result_column!r}")
            for row in range(2, sheet.max_row + 1):
                raw_id = sheet.cell(row, headers["ID"]).value
                if raw_id is None or str(raw_id).strip() == "":
                    if any(sheet.cell(row, headers[name]).value not in (None, "") for name in required[1:]):
                        raise BridgeError(f"Criterion row {sheet.title}!{row} has no ID")
                    continue
                criterion_id = str(raw_id).strip()
                if criterion_id in seen_ids:
                    raise BridgeError(f"Duplicate criterion ID: {criterion_id}")
                seen_ids.add(criterion_id)
                criteria.append({
                    "criterion_key": f"{sheet.title}!{row}:{criterion_id}",
                    "criterion_id": criterion_id,
                    "sheet": sheet.title,
                    "row": row,
                    "group": sheet.cell(row, headers["Nhóm"]).value,
                    "criterion": sheet.cell(row, headers["Tiêu chí"]).value,
                    "requirement": sheet.cell(row, headers["Requirement"]).value,
                    "analysis_instruction": sheet.cell(row, headers["Analysis Instruction"]).value,
                })
    finally:
        workbook.close()
    if criteria_sheets != 1:
        raise BridgeError(f"Expected exactly one criteria sheet, found {criteria_sheets}")
    if len(criteria) != paths.expected_criteria_count:
        raise BridgeError(f"Expected {paths.expected_criteria_count} criteria, found {len(criteria)}")
    payload = {
        "schema_version": SCHEMA_VERSION,
        "source": {
            "content_a": _relative(paths.content_a, paths.project_root),
            "content_b": _relative(paths.content_b, paths.project_root),
            "criteria": _relative(paths.criteria, paths.project_root),
            "content_a_sha256": _sha256(paths.content_a),
            "content_b_sha256": _sha256(paths.content_b),
            "criteria_sha256": _sha256(paths.criteria),
            "result_column": paths.result_column,
            "expected_criteria_count": paths.expected_criteria_count,
        },
        "parts": parts,
        "criteria": criteria,
    }
    payload["payload_digest"] = payload_digest(payload)
    return payload


def _write_json_atomic(path: Path, data: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(data, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        os.replace(name, path)
    except Exception:
        Path(name).unlink(missing_ok=True)
        raise


def prepare(paths: Paths) -> Dict[str, Any]:
    payload = extract_criteria(paths)
    _write_json_atomic(paths.criteria_payload, payload)
    return payload


def _load_json(path: Path) -> Any:
    def reject_duplicates(pairs: List[Tuple[str, Any]]) -> Dict[str, Any]:
        document: Dict[str, Any] = {}
        for key, value in pairs:
            if key in document:
                raise BridgeError(f"Duplicate JSON field {key!r} in {path}")
            document[key] = value
        return document

    try:
        return json.loads(path.read_text(encoding="utf-8-sig"), object_pairs_hook=reject_duplicates)
    except json.JSONDecodeError as error:
        raise BridgeError(f"Invalid JSON in {path}: {error}") from error


def validate_results(payload: Mapping[str, Any], document: Any) -> Dict[str, Dict[str, str]]:
    if not isinstance(document, dict) or set(document) != {"schema_version", "payload_digest", "parts"}:
        raise BridgeError("Result document must contain only schema_version, payload_digest and parts")
    if document["schema_version"] != SCHEMA_VERSION:
        raise BridgeError(f"Unsupported schema_version: {document['schema_version']!r}")
    if document["payload_digest"] != payload["payload_digest"]:
        raise BridgeError("Results refer to a different payload; analyze the current prepared inputs")
    if not isinstance(document["parts"], list):
        raise BridgeError("parts must be an array")
    expected_parts = [part["part_id"] for part in payload["parts"]]
    actual_parts = [part.get("part_id") if isinstance(part, dict) else None for part in document["parts"]]
    if actual_parts != expected_parts:
        raise BridgeError("Missing, unknown, duplicate or out-of-order Part results")
    expected_keys = [item["criterion_key"] for item in payload["criteria"]]
    mapped: Dict[str, Dict[str, str]] = {}
    for part in document["parts"]:
        if set(part) != {"part_id", "results"} or not isinstance(part["results"], list):
            raise BridgeError(f"Invalid fields for {part['part_id']}")
        items = part["results"]
        actual_keys = [item.get("criterion_key") if isinstance(item, dict) else None for item in items]
        if actual_keys != expected_keys:
            raise BridgeError(f"{part['part_id']}: missing, unknown, duplicate or out-of-order criteria")
        part_results: Dict[str, str] = {}
        for item, criterion in zip(items, payload["criteria"]):
            if set(item) != {"criterion_key", "criterion_id", "result"}:
                raise BridgeError(f"{part['part_id']}: invalid criterion fields")
            if item["criterion_id"] != criterion["criterion_id"]:
                raise BridgeError(f"{part['part_id']}: criterion_id does not match {criterion['criterion_key']}")
            if not isinstance(item["result"], str):
                raise BridgeError(f"{part['part_id']}: Result must be a string for {criterion['criterion_id']}")
            result = item["result"].strip()
            if len(result) > 32767:
                raise BridgeError(f"{part['part_id']}: Result exceeds Excel cell limit for {criterion['criterion_id']}")
            part_results[item["criterion_key"]] = result
        mapped[part["part_id"]] = part_results
    return mapped


def finalize(paths: Paths, overwrite: bool = False) -> int:
    _require_files((paths.content_a, paths.content_b, paths.criteria, paths.criteria_payload, paths.analysis_results))
    payload = _load_json(paths.criteria_payload)
    if payload != extract_criteria(paths):
        raise BridgeError("criteria_payload is stale; run prepare again")
    mapped = validate_results(payload, _load_json(paths.analysis_results))
    if paths.output in (paths.content_a, paths.content_b, paths.criteria, paths.criteria_payload, paths.analysis_results):
        raise BridgeError("Output must not overwrite an input file")
    if paths.output.exists() and not overwrite:
        raise BridgeError("Output already exists; choose a new output path or use --overwrite for an intentional replacement")
    paths.output.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=f".{paths.output.stem}.", suffix=".xlsx", dir=paths.output.parent)
    os.close(fd)
    temporary = Path(name)
    expected_rows: Dict[str, List[List[Any]]] = {}
    try:
        workbook = Workbook()
        try:
            workbook.remove(workbook.active)
            for part in payload["parts"]:
                part_id = part["part_id"]
                sheet = workbook.create_sheet(part_id)
                rows = [OUTPUT_HEADERS.copy()]
                for criterion in payload["criteria"]:
                    result = mapped[part_id][criterion["criterion_key"]]
                    rows.append([
                        criterion["criterion_id"], criterion["group"], criterion["criterion"],
                        criterion["requirement"], result or None,
                    ])
                expected_rows[part_id] = rows
                for row in rows:
                    sheet.append(row)
                    # Treat generated text as text, including a leading '='.
                    for cell in sheet[sheet.max_row]:
                        if isinstance(cell.value, str):
                            cell.data_type = "s"
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="244062")
                for column, width in zip("ABCDE", (12, 24, 36, 65, 100)):
                    sheet.column_dimensions[column].width = width
                sheet.freeze_panes = "E2"
                sheet.auto_filter.ref = sheet.dimensions
            workbook.save(temporary)
        finally:
            workbook.close()
        output = load_workbook(temporary, data_only=False)
        try:
            if output.sheetnames != list(expected_rows):
                raise BridgeError("Output verification failed: unexpected sheets")
            for part_id, rows in expected_rows.items():
                actual = [list(row) for row in output[part_id].iter_rows(values_only=True)]
                if actual != rows:
                    raise BridgeError(f"Output verification failed: data mismatch in {part_id}")
        finally:
            output.close()
        os.replace(temporary, paths.output)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return len(payload["parts"]) * len(payload["criteria"])


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("prepare", "finalize"))
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("--project-root", type=Path, default=Path.cwd())
    parser.add_argument("--overwrite", action="store_true", help="Intentionally replace an existing output workbook")
    args = parser.parse_args(argv)
    try:
        paths = read_config(args.config, args.project_root)
        if args.command == "prepare":
            payload = prepare(paths)
            print(f"Prepared {len(payload['parts'])} Parts x {len(payload['criteria'])} criteria: {paths.criteria_payload}")
        else:
            count = finalize(paths, overwrite=args.overwrite)
            print(f"Wrote {count} criterion rows in Part sheets: {paths.output}")
        return 0
    except (BridgeError, OSError, KeyError, ValueError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
