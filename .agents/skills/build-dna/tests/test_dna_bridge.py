"""I/O fixtures only: simulated decisions are not learned author DNA."""
import copy
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "dna_bridge.py"
spec = importlib.util.spec_from_file_location("dna_bridge", SCRIPT)
bridge = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = bridge
spec.loader.exec_module(bridge)


class BridgeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="dna-test-")
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        self.paths = bridge.Paths(root / "result_analysis.xlsx", root / "grouped.json",
                                  root / "decisions.json", root / "AUTHOR_DNA.xlsx")
        workbook = Workbook()
        workbook.remove(workbook.active)
        # Deliberately create sheets and rows out of order; IDs, not row indices, map observations.
        for name in ("Part10", "Part2", "Part1"):
            sheet = workbook.create_sheet(name)
            sheet.append(bridge.INPUT_HEADERS)
            ids = ("C460", "C100", "C999") if name == "Part2" else ("C100", "C460", "C999")
            for criterion_id in ids:
                result = f"{name}: giữ thông tin chính, nén so sánh phụ ({criterion_id})."
                if criterion_id == "C999":
                    result = None
                sheet.append([criterion_id, "Nhóm thử", f"Tiêu chí {criterion_id}", "Đối chiếu A→B", result])
        workbook.create_sheet("Notes").append(["Ignored"])
        workbook.save(self.paths.input)
        workbook.close()
        self.original = self.paths.input.read_bytes()
        self.payload = bridge.prepare(self.paths)

    def decisions(self):
        items = []
        for criterion in self.payload["criteria"]:
            item = {"criterion_id": criterion["criterion_id"], "decision": "drop",
                    "synthesis": "Không đủ căn cứ cho một rule độc lập.", "confidence": "Low",
                    "evidence": [], "filter_reason": "Thiếu evidence xuyên đối tượng."}
            item.update({field: "" for field in bridge.EXECUTION_FIELDS})
            if criterion["criterion_id"] == "C100":
                item.update(decision="keep", synthesis="Giữ thông tin chính, nén phần so sánh phụ.",
                            confidence="High", filter_reason="", when="Khi source có thông tin chính và so sánh phụ.",
                            action="Giữ mô tả đối tượng chính, nén so sánh phụ.",
                            selection="Ưu tiên thông tin trực tiếp mô tả đối tượng.",
                            validation="Đối chiếu output còn mô tả chính và không lặp so sánh phụ.",
                            condition_variation="Chưa có evidence về variation; chỉ áp dụng trong phạm vi WHEN.",
                            target_step="STEP 1")
                item["evidence"] = [{"part_id": obs["part_id"], "result_excerpt": "giữ thông tin chính, nén so sánh phụ",
                                     "role": "support"} for obs in criterion["observations"][:2]]
            items.append(item)
        return {"schema_version": "1.0", "payload_digest": self.payload["payload_digest"], "criteria": items}

    def save_decisions(self, document=None):
        bridge.write_json(self.paths.dna_results, document or self.decisions())

    def edit_input(self, edit):
        workbook = load_workbook(self.paths.input)
        try:
            edit(workbook)
            workbook.save(self.paths.input)
        finally:
            workbook.close()

    def test_prepare_preserves_results_blanks_and_coordinates(self):
        self.assertEqual(self.payload["parts"], ["Part1", "Part2", "Part10"])
        self.assertEqual(self.payload["ignored_sheets"], ["Notes"])
        self.assertEqual([c["criterion_id"] for c in self.payload["criteria"]], ["C100", "C460", "C999"])
        obs = self.payload["criteria"][0]["observations"]
        self.assertEqual(obs[1]["result_cell"], "Part2!E3")
        self.assertEqual(obs[1]["result"], "Part2: giữ thông tin chính, nén so sánh phụ (C100).")
        self.assertEqual([o["result"] for o in self.payload["criteria"][2]["observations"]], ["", "", ""])
        self.assertEqual(self.paths.input.read_bytes(), self.original)
        self.assertNotIn("decision", self.payload["criteria"][0])

    def test_export_keeps_only_selected_id_and_four_columns(self):
        document = self.decisions()
        document["criteria"][0]["action"] = "=Keep as literal text"
        self.save_decisions(document)
        self.assertEqual(bridge.finalize(self.paths), 1)
        workbook = load_workbook(self.paths.output)
        try:
            self.assertEqual(workbook.sheetnames, ["AUTHOR_DNA"])
            sheet = workbook.active
            self.assertEqual([c.value for c in sheet[1]], bridge.OUTPUT_HEADERS)
            self.assertEqual((sheet.max_row, sheet.max_column), (2, 4))
            self.assertEqual(sheet["A2"].value, "C100")
            self.assertIn("ACTION\n=Keep as literal text", sheet["C2"].value)
            self.assertEqual(sheet["C2"].data_type, "s")
            self.assertEqual(sheet.freeze_panes, "C2")
        finally:
            workbook.close()
        self.assertEqual(self.paths.input.read_bytes(), self.original)

    def test_all_dropped_exports_header_only(self):
        document = self.decisions()
        item = document["criteria"][0]
        item.update(decision="drop", filter_reason="Chưa đủ bằng chứng.", confidence="Low")
        item.update({field: "" for field in bridge.EXECUTION_FIELDS})
        self.save_decisions(document)
        self.assertEqual(bridge.finalize(self.paths), 0)
        workbook = load_workbook(self.paths.output)
        try:
            self.assertEqual(workbook.active.max_row, 1)
            self.assertEqual([c.value for c in workbook.active[1]], bridge.OUTPUT_HEADERS)
        finally:
            workbook.close()

    def test_invalid_input_rejected_without_replacing_payload(self):
        original_payload = self.paths.grouped_payload.read_bytes()
        edits = [
            lambda w: w["Part2"].delete_rows(2),
            lambda w: setattr(w["Part2"]["A3"], "value", "C460"),
            lambda w: setattr(w["Part2"]["D3"], "value", "Different Requirement"),
            lambda w: setattr(w["Part2"]["E3"], "value", "=1+1"),
            lambda w: setattr(w["Part2"]["E3"], "value", 12),
            lambda w: setattr(w["Part2"]["A3"], "value", "C777"),
            lambda w: setattr(w["Part2"]["A1"], "value", "Criterion ID"),
            lambda w: setattr(w["Part2"], "title", "Part02"),
        ]
        for edit in edits:
            with self.subTest(edit=edit):
                self.paths.input.write_bytes(self.original)
                self.edit_input(edit)
                with self.assertRaises(bridge.BridgeError):
                    bridge.prepare(self.paths)
                self.assertEqual(self.paths.grouped_payload.read_bytes(), original_payload)

    def test_single_part_cannot_support_dna(self):
        document = self.decisions()
        evidence = document["criteria"][0]["evidence"]
        evidence[1] = copy.deepcopy(evidence[0])
        with self.assertRaises(bridge.BridgeError):
            bridge.validate_results(self.payload, document)

    def test_empty_or_foreign_evidence_rejected(self):
        for change in ("invented quote", "", "   "):
            document = self.decisions()
            document["criteria"][0]["evidence"][0]["result_excerpt"] = change
            with self.subTest(quote=change), self.assertRaises(bridge.BridgeError):
                bridge.validate_results(self.payload, document)
        document = self.decisions()
        document["criteria"][0]["evidence"][0]["part_id"] = "Part99"
        with self.assertRaises(bridge.BridgeError):
            bridge.validate_results(self.payload, document)

    def test_low_confidence_or_unresolved_conflict_cannot_be_kept(self):
        for change in ("Low", "unresolved"):
            document = self.decisions()
            item = document["criteria"][0]
            if change == "Low":
                item["confidence"] = change
            else:
                item["evidence"].append({"part_id": "Part10", "result_excerpt": "giữ thông tin chính", "role": change})
            with self.subTest(change=change), self.assertRaises(bridge.BridgeError):
                bridge.validate_results(self.payload, document)

    def test_every_id_requires_exactly_one_decision(self):
        for mode in ("missing", "duplicate", "order", "unknown"):
            document = self.decisions()
            items = document["criteria"]
            if mode == "missing":
                items.pop()
            elif mode == "duplicate":
                items.append(copy.deepcopy(items[0]))
            elif mode == "order":
                items.reverse()
            else:
                items[0]["criterion_id"] = "C000"
            with self.subTest(mode=mode), self.assertRaises(bridge.BridgeError):
                bridge.validate_results(self.payload, document)

    def test_stale_input_and_payload_rejected(self):
        self.save_decisions()
        self.edit_input(lambda w: setattr(w["Part1"]["E2"], "value", "Changed observation"))
        with self.assertRaisesRegex(bridge.BridgeError, "stale"):
            bridge.finalize(self.paths)
        self.paths.input.write_bytes(self.original)
        modified = copy.deepcopy(self.payload)
        modified["criteria"][0]["observations"][0]["result"] = "Edited payload"
        bridge.write_json(self.paths.grouped_payload, modified)
        with self.assertRaisesRegex(bridge.BridgeError, "stale"):
            bridge.finalize(self.paths)

    def test_existing_output_preserved_on_failure_and_explicit_overwrite(self):
        self.save_decisions()
        bridge.finalize(self.paths)
        old = self.paths.output.read_bytes()
        with self.assertRaisesRegex(bridge.BridgeError, "Output exists"):
            bridge.finalize(self.paths)
        self.assertEqual(self.paths.output.read_bytes(), old)
        invalid = self.decisions()
        invalid["criteria"][0]["validation"] = ""
        self.save_decisions(invalid)
        with self.assertRaises(bridge.BridgeError):
            bridge.finalize(self.paths, overwrite=True)
        self.assertEqual(self.paths.output.read_bytes(), old)
        self.save_decisions()
        self.assertEqual(bridge.finalize(self.paths, overwrite=True), 1)

    def test_excel_limit_prevents_silent_truncation(self):
        document = self.decisions()
        document["criteria"][0]["action"] = "x" * 32767
        self.save_decisions(document)
        with self.assertRaisesRegex(bridge.BridgeError, "cell limits"):
            bridge.finalize(self.paths)
        self.assertFalse(self.paths.output.exists())

    def test_config_prevents_input_overwrite_and_resolves_project_root(self):
        config = self.paths.input.parent / "config.yaml"
        config.write_text("input: result_analysis.xlsx\ngrouped_payload: grouped.json\ndna_results: decisions.json\noutput: AUTHOR_DNA.xlsx\n", encoding="utf-8")
        self.assertEqual(bridge.read_config(config, config.parent), self.paths)
        config.write_text(config.read_text(encoding="utf-8").replace("AUTHOR_DNA.xlsx", "result_analysis.xlsx"), encoding="utf-8")
        with self.assertRaises(bridge.BridgeError):
            bridge.read_config(config, config.parent)

    def test_duplicate_json_keys_rejected(self):
        self.paths.dna_results.write_text('{"criteria": [], "criteria": []}', encoding="utf-8")
        with self.assertRaises(bridge.BridgeError):
            bridge.load_json(self.paths.dna_results)

    def test_payload_digest_links_decisions_to_current_input(self):
        document = self.decisions()
        document["payload_digest"] = "0" * 64
        with self.assertRaises(bridge.BridgeError):
            bridge.validate_results(self.payload, document)


if __name__ == "__main__":
    unittest.main()
