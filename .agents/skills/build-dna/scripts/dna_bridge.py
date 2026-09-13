#!/usr/bin/env python3
"""Group Part observations and export LLM-authored DNA; no semantic synthesis."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

VERSION = "1.0"
INPUT_HEADERS = ["ID", "Nhóm", "Tiêu chí", "Requirement", "Result"]
OUTPUT_HEADERS = ["ID", "Tiêu chí", "DNA Rule", "Condition/Variation"]
EXECUTION_FIELDS = ("when", "action", "selection", "validation", "condition_variation", "target_step")
DECISION_FIELDS = {"criterion_id", "decision", "synthesis", "confidence", "evidence", "filter_reason", *EXECUTION_FIELDS}
PART = re.compile(r"Part([1-9][0-9]*)\Z")
CRITERION = re.compile(r"C[0-9]+\Z")


class BridgeError(RuntimeError):
    pass


@dataclass(frozen=True)
class Paths:
    input: Path
    grouped_payload: Path
    dna_results: Path
    output: Path


def read_config(config: Path, root: Path) -> Paths:
    # Same flat key: value convention as Skill 1; no YAML dependency required.
    values = {}
    for number, raw in enumerate(config.read_text(encoding="utf-8-sig").splitlines(), 1):
        if not raw.strip() or raw.lstrip().startswith("#"):
            continue
        if ":" not in raw:
            raise BridgeError(f"Invalid config line {number}")
        key, value = (piece.strip() for piece in raw.split(":", 1))
        if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'":
            value = value[1:-1]
        if key in values or not value:
            raise BridgeError(f"Duplicate key or empty value at config line {number}")
        values[key] = value
    if set(values) != {"input", "grouped_payload", "dna_results", "output"}:
        raise BridgeError("Config requires only input, grouped_payload, dna_results, output")
    paths = Paths(**{key: (root / value).resolve() for key, value in values.items()})
    validate_paths(paths)
    return paths


def validate_paths(paths: Paths) -> None:
    resolved = [path.resolve() for path in vars(paths).values()]
    if len(set(resolved)) != len(resolved):
        raise BridgeError("Input and artifact paths must all be distinct")
    existing = [path for path in resolved if path.exists()]
    for i, path in enumerate(existing):
        if any(os.path.samefile(path, other) for other in existing[i + 1:]):
            raise BridgeError("Input and artifact paths must not alias the same file")
    if paths.input.suffix.lower() != ".xlsx" or paths.output.suffix.lower() != ".xlsx":
        raise BridgeError("Input and output must be .xlsx files")
    if any(path.suffix.lower() != ".json" for path in (paths.grouped_payload, paths.dna_results)):
        raise BridgeError("Payload and decisions must be .json files")


def digest(payload: dict) -> str:
    body = {key: value for key, value in payload.items() if key != "payload_digest"}
    return hashlib.sha256(json.dumps(body, ensure_ascii=False, sort_keys=True,
                                     separators=(",", ":")).encode("utf-8")).hexdigest()


def extract(paths: Paths) -> dict:
    validate_paths(paths)
    source_hash = hashlib.sha256(paths.input.read_bytes()).hexdigest()
    workbook = load_workbook(paths.input, read_only=True, data_only=False)
    try:
        sheets = sorted((sheet for sheet in workbook if PART.fullmatch(sheet.title)),
                        key=lambda sheet: int(PART.fullmatch(sheet.title).group(1)))
        if not sheets:
            raise BridgeError("Input has no Part sheets (Part1, Part2, ...)")
        ignored = [sheet.title for sheet in workbook if not PART.fullmatch(sheet.title)]
        malformed = [name for name in ignored if name.lower().startswith("part")]
        if malformed:
            raise BridgeError(f"Invalid Part sheet names: {malformed}")
        grouped: dict[str, dict] = {}
        part_ids = []
        for sheet in sheets:
            part_ids.append(sheet.title)
            rows = sheet.iter_rows()
            headers = [cell.value for cell in next(rows, [])]
            if headers != INPUT_HEADERS:
                raise BridgeError(f"{sheet.title}: expected exactly {INPUT_HEADERS}")
            seen = set()
            for row in rows:
                values = [cell.value for cell in row]
                if all(value in (None, "") for value in values):
                    continue
                if any(cell.data_type in ("f", "e") for cell in row):
                    raise BridgeError(f"{sheet.title}!{row[0].row}: formulas/errors are not observations")
                if any(value is not None and not isinstance(value, str) for value in values):
                    raise BridgeError(f"{sheet.title}!{row[0].row}: criterion fields must be text")
                criterion_id, group, criterion, requirement, result = values
                if not isinstance(criterion_id, str) or not CRITERION.fullmatch(criterion_id):
                    raise BridgeError(f"{sheet.title}!{row[0].row}: invalid criterion ID")
                if criterion_id in seen:
                    raise BridgeError(f"{sheet.title}: duplicate ID {criterion_id}")
                if not criterion or not criterion.strip() or not requirement or not requirement.strip():
                    raise BridgeError(f"{sheet.title}: missing criterion name or Requirement for {criterion_id}")
                seen.add(criterion_id)
                metadata = {"criterion_id": criterion_id, "group": group,
                            "criterion": criterion, "requirement": requirement}
                if len(part_ids) == 1:
                    grouped[criterion_id] = {**metadata, "observations": []}
                elif criterion_id not in grouped:
                    raise BridgeError(f"{sheet.title}: unknown criterion {criterion_id}")
                elif any(grouped[criterion_id][key] != value for key, value in metadata.items()):
                    raise BridgeError(f"{sheet.title}: inconsistent metadata for {criterion_id}")
                grouped[criterion_id]["observations"].append({
                    "part_id": sheet.title, "row": row[0].row,
                    "result_cell": f"{sheet.title}!E{row[0].row}",
                    "result": result if result is not None else "",
                })
            if not seen:
                raise BridgeError(f"{sheet.title}: no criteria")
            if seen != set(grouped):
                raise BridgeError(f"{sheet.title}: missing criteria {sorted(set(grouped) - seen)}")
        payload = {"schema_version": VERSION,
                   "source": {"input": str(paths.input.resolve()), "sha256": source_hash},
                   "parts": part_ids, "ignored_sheets": ignored,
                   "criteria": list(grouped.values())}
        payload["payload_digest"] = digest(payload)
        return payload
    finally:
        workbook.close()


def write_json(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(data, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)


def load_json(path: Path) -> Any:
    def unique(pairs):
        value = {}
        for key, item in pairs:
            if key in value:
                raise BridgeError(f"Duplicate JSON key {key!r} in {path}")
            value[key] = item
        return value
    return json.loads(path.read_text(encoding="utf-8-sig"), object_pairs_hook=unique)


def prepare(paths: Paths) -> dict:
    payload = extract(paths)
    write_json(paths.grouped_payload, payload)
    return payload


def validate_results(payload: dict, document: Any) -> list[dict]:
    if not isinstance(document, dict) or set(document) != {"schema_version", "payload_digest", "criteria"}:
        raise BridgeError("Decisions require only schema_version, payload_digest, criteria")
    if document["schema_version"] != VERSION or document["payload_digest"] != payload["payload_digest"]:
        raise BridgeError("Decisions refer to a different prepared payload")
    decisions = document["criteria"]
    if not isinstance(decisions, list):
        raise BridgeError("criteria must be an array")
    actual_ids = [item.get("criterion_id") if isinstance(item, dict) else None for item in decisions]
    if actual_ids != [item["criterion_id"] for item in payload["criteria"]]:
        raise BridgeError("Missing, duplicate, unknown or out-of-order criterion decisions")
    kept = []
    for decision, criterion in zip(decisions, payload["criteria"]):
        label = criterion["criterion_id"]
        if set(decision) != DECISION_FIELDS:
            raise BridgeError(f"{label}: invalid decision fields")
        for field in DECISION_FIELDS - {"evidence"}:
            if not isinstance(decision[field], str):
                raise BridgeError(f"{label}: {field} must be a string")
        if decision["decision"] not in ("keep", "drop") or decision["confidence"] not in ("High", "Medium", "Low"):
            raise BridgeError(f"{label}: invalid decision or confidence")
        if not decision["synthesis"].strip() or not isinstance(decision["evidence"], list):
            raise BridgeError(f"{label}: synthesis required and evidence must be an array")
        observations = {item["part_id"]: item["result"] for item in criterion["observations"]}
        supporting_parts = set()
        unresolved = False
        for evidence in decision["evidence"]:
            if not isinstance(evidence, dict) or set(evidence) != {"part_id", "result_excerpt", "role"}:
                raise BridgeError(f"{label}: invalid evidence fields")
            if any(not isinstance(value, str) for value in evidence.values()):
                raise BridgeError(f"{label}: evidence fields must be strings")
            part_id, quote, role = evidence["part_id"], evidence["result_excerpt"], evidence["role"]
            if part_id not in observations or not quote.strip() or quote not in observations[part_id]:
                raise BridgeError(f"{label}: excerpt not found in {part_id}.Result")
            if role not in ("support", "variation", "exception", "unresolved"):
                raise BridgeError(f"{label}: invalid evidence role")
            if role in ("support", "variation"):
                supporting_parts.add(part_id)
            unresolved |= role == "unresolved"
        if decision["decision"] == "drop":
            if not decision["filter_reason"].strip() or any(decision[field] != "" for field in EXECUTION_FIELDS):
                raise BridgeError(f"{label}: dropped IDs need a reason and empty execution fields")
            continue
        # Validate the LLM decision, never change it or synthesize a substitute.
        if decision["confidence"] == "Low" or unresolved or len(supporting_parts) < 2:
            raise BridgeError(f"{label}: kept ID needs High/Medium, two distinct evidence Parts, no unresolved conflict")
        if decision["filter_reason"] != "" or any(not decision[field].strip() for field in EXECUTION_FIELDS):
            raise BridgeError(f"{label}: kept ID needs complete execution fields and empty filter_reason")
        if decision["target_step"] not in ("STEP 1", "STEP 2", "STEP 3", "STEP 4"):
            raise BridgeError(f"{label}: invalid target_step")
        kept.append(decision)
    return kept


def output_rows(payload: dict, decisions: list[dict]) -> list[list[str]]:
    names = {item["criterion_id"]: item["criterion"] for item in payload["criteria"]}
    rows = [OUTPUT_HEADERS.copy()]
    for item in decisions:
        rule = "\n\n".join(f"{label}\n{item[field]}" for label, field in (
            ("WHEN", "when"), ("ACTION", "action"), ("SELECTION", "selection"), ("VALIDATION", "validation")))
        row = [item["criterion_id"], names[item["criterion_id"]], rule, item["condition_variation"]]
        if any(len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value) for value in row):
            raise BridgeError(f"{item['criterion_id']}: output exceeds Excel cell limits or contains illegal characters")
        rows.append(row)
    return rows


def finalize(paths: Paths, overwrite: bool = False) -> int:
    payload = load_json(paths.grouped_payload)
    if payload != extract(paths):
        raise BridgeError("Prepared payload is stale or modified; prepare and synthesize again")
    kept = validate_results(payload, load_json(paths.dna_results))
    rows = output_rows(payload, kept)
    if paths.output.exists() and not overwrite:
        raise BridgeError("Output exists; choose a new path or --overwrite for intentional replacement")
    paths.output.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=f".{paths.output.stem}.", suffix=".xlsx", dir=paths.output.parent)
    os.close(fd)
    temporary = Path(name)
    try:
        workbook = Workbook()
        try:
            sheet = workbook.active
            sheet.title = "AUTHOR_DNA"
            for row in rows:
                sheet.append(row)
                for cell in sheet[sheet.max_row]:
                    cell.data_type = "s"
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="244062")
            for column, width in zip("ABCD", (12, 36, 95, 75)):
                sheet.column_dimensions[column].width = width
            sheet.freeze_panes = "C2"
            sheet.auto_filter.ref = sheet.dimensions
            workbook.save(temporary)
        finally:
            workbook.close()
        check = load_workbook(temporary, read_only=True, data_only=False)
        try:
            if check.sheetnames != ["AUTHOR_DNA"] or [list(row) for row in check.active.iter_rows(values_only=True)] != rows:
                raise BridgeError("Saved workbook failed round-trip verification")
        finally:
            check.close()
        # Do not replace a prior result if the input changed during export.
        if hashlib.sha256(paths.input.read_bytes()).hexdigest() != payload["source"]["sha256"]:
            raise BridgeError("Input changed during export; prepare and synthesize again")
        os.replace(temporary, paths.output)
    finally:
        temporary.unlink(missing_ok=True)
    return len(kept)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("prepare", "finalize"))
    parser.add_argument("--config", type=Path, required=True)
    parser.add_argument("--project-root", type=Path, default=Path.cwd())
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args(argv)
    try:
        paths = read_config(args.config, args.project_root)
        if args.command == "prepare":
            payload = prepare(paths)
            print(f"Grouped {len(payload['parts'])} Parts x {len(payload['criteria'])} criteria: {paths.grouped_payload}")
            if payload["ignored_sheets"]:
                print(f"Ignored non-Part sheets: {payload['ignored_sheets']}")
        else:
            count = finalize(paths, args.overwrite)
            print(f"Exported {count} DNA criteria: {paths.output}")
            if count == 0:
                print("No DNA rules had sufficient evidence; workbook contains headers only.")
        return 0
    except (BridgeError, OSError, ValueError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
